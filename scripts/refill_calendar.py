#!/usr/bin/env python3
"""
Refill the Calendar (90-day) tab with Spoke rows drawn from unused
Master Keywords.

Why this exists
---------------
`generate-post.py` only publishes Calendar rows whose Type contains
"spoke". Pillar-update and New-pillar rows are deliberately skipped as
human-authoring work. When the last Spoke row is published the runner
finds nothing eligible, exits clean, and commits nothing. The pipeline
looks healthy and silently stops publishing. That is exactly what
happened between 2026-08-19 and 2026-08-31.

This script tops the Calendar back up. It is idempotent: a keyword that
already appears anywhere in the Calendar tab is never added twice, so it
is safe to re-run.

Usage:
    python scripts/refill_calendar.py            # append missing rows
    python scripts/refill_calendar.py --dry-run  # show what would change
"""

import sys
from datetime import date, timedelta
from pathlib import Path

from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parent.parent
WORKBOOK_PATH = REPO_ROOT / "GrowthMax-Keyword-Program.xlsx"

# Calendar columns (1-indexed):
#   A Week | B Publish date | C Pillar | D Type | E Working title
#   F Primary keyword | G Owner | H Status
COL_WEEK, COL_DATE, COL_PILLAR, COL_TYPE = 1, 2, 3, 4
COL_TITLE, COL_KEYWORD, COL_OWNER, COL_STATUS = 5, 6, 7, 8
DATA_START_ROW = 5

# Posting cadence: Monday, Wednesday, Friday (matches the cron in
# .github/workflows/generate-post.yml).
PUBLISH_WEEKDAYS = {0, 2, 4}

# Working titles are deliberately distinct from every published slug so the
# early slug-collision guard in generate-post.py does not fire. Keywords must
# match the Master Keywords tab exactly, because writeback_publish() syncs
# status back by exact keyword match.
QUEUE = [
    # (pillar, keyword, working title)
    ("2", "custom AI agents for enterprise",
     "What Makes an AI Agent Enterprise Ready"),
    ("4", "enterprise AI training",
     "What Effective Enterprise AI Training Actually Looks Like"),
    ("2", "enterprise AI agent development",
     "How Enterprise AI Agent Development Actually Works"),
    ("1", "AI maturity model",
     "An AI Maturity Model That Skips the Hype"),
    ("4", "AI literacy program",
     "How to Build an AI Literacy Program That Sticks"),
    ("2", "AI agent integration",
     "Connecting AI Agents to the Systems You Already Run"),
    ("5", "human AI collaboration",
     "Designing Human and AI Collaboration That Holds Up"),
    ("3", "AI implementation timeline",
     "How Long AI Implementation Really Takes"),
    ("2", "AI agent evaluation",
     "How to Evaluate an AI Agent Before You Trust It"),
    ("4", "AI skills for employees",
     "The AI Skills Your Employees Actually Need"),
    ("1", "AI investment prioritization",
     "How to Prioritize AI Investments When Everything Looks Urgent"),
    ("2", "AI agent observability",
     "Knowing What Your AI Agent Did, and Why"),
    ("3", "AI adoption metrics KPIs",
     "The AI Adoption KPIs Worth Reporting to Your Board"),
    ("5", "responsible AI enterprise",
     "Responsible AI in the Enterprise, Past the Policy Document"),
    ("4", "AI bootcamp for engineers",
     "What an AI Bootcamp for Engineers Should Cover"),
    ("2", "AI agents for HR",
     "Where AI Agents Fit in HR Without Replacing the Human Part"),
    ("4", "custom AI training for companies",
     "When Off-the-Shelf AI Training Is Not Enough"),
]


def next_publish_dates(start, count):
    """Return `count` Mon/Wed/Fri dates on or after `start`."""
    out, cursor = [], start
    while len(out) < count:
        if cursor.weekday() in PUBLISH_WEEKDAYS:
            out.append(cursor)
        cursor += timedelta(days=1)
    return out


def existing_keywords(sheet):
    """Every keyword already present in the Calendar, lowercased."""
    found = set()
    for row_idx in range(DATA_START_ROW, sheet.max_row + 1):
        kw = sheet.cell(row=row_idx, column=COL_KEYWORD).value
        if kw:
            found.add(str(kw).strip().lower())
    return found


def latest_calendar_date(sheet):
    """Latest publish date on any Calendar row, or None."""
    dates = []
    for row_idx in range(DATA_START_ROW, sheet.max_row + 1):
        val = sheet.cell(row=row_idx, column=COL_DATE).value
        if hasattr(val, "date"):
            val = val.date()
        if isinstance(val, date):
            dates.append(val)
    return max(dates) if dates else None


def main():
    dry_run = "--dry-run" in sys.argv

    wb = load_workbook(WORKBOOK_PATH)
    cal = wb["Calendar (90-day)"]

    already = existing_keywords(cal)
    pending = [item for item in QUEUE if item[1].strip().lower() not in already]

    if not pending:
        print("Calendar already contains every queued keyword. Nothing to do.")
        return 0

    # Start scheduling the day after the latest existing Calendar date, but
    # never in the past. Overdue rows are picked first by the runner, so a
    # small backlog is fine and lets the pipeline catch up.
    anchor = max(date.today(), (latest_calendar_date(cal) or date.today()))
    dates = next_publish_dates(anchor + timedelta(days=1), len(pending))

    write_row = cal.max_row + 1
    print("Appending %d Spoke rows starting at row %d:\n" % (len(pending), write_row))

    for (pillar, keyword, title), publish_on in zip(pending, dates):
        print("  %s  P%s  %s" % (publish_on, pillar, title))
        print("              kw: %s" % keyword)
        if not dry_run:
            cal.cell(row=write_row, column=COL_WEEK).value = "(refill 2026-08-31)"
            cal.cell(row=write_row, column=COL_DATE).value = publish_on
            cal.cell(row=write_row, column=COL_PILLAR).value = int(pillar)
            cal.cell(row=write_row, column=COL_TYPE).value = "Spoke"
            cal.cell(row=write_row, column=COL_TITLE).value = title
            cal.cell(row=write_row, column=COL_KEYWORD).value = keyword
            cal.cell(row=write_row, column=COL_OWNER).value = "auto"
            cal.cell(row=write_row, column=COL_STATUS).value = "Not started"
        write_row += 1

    if dry_run:
        print("\n[dry run] Workbook not modified.")
        return 0

    wb.save(WORKBOOK_PATH)
    print("\nWorkbook saved. Runway extended through %s." % dates[-1])
    return 0


if __name__ == "__main__":
    sys.exit(main())
